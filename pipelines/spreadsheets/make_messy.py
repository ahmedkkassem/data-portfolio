"""Build the "before" files for the spreadsheet-rescue demo.

Takes the real November 2011 sales lines from the warehouse and writes them out the way three sales
teams might keep them: each team with its own layout and habits, plus the classic spreadsheet problems
(title rows, subtotal rows, pasted-twice blocks, text dates, decimal commas, "12 pcs", local country
names, flipped return signs). The pipeline in clean.py only ever sees these files.

Deterministic (fixed seed), so the same files are rebuilt on every run. Also writes small preview JSON
for the web page, including where each problem was planted.
"""
import csv
import datetime as dt
import json
import pathlib
import random

import duckdb
from openpyxl import Workbook

OUT = pathlib.Path("build/site/spreadsheets/files")
PREVIEW = pathlib.Path("build/site/data/spreadsheets_before.json")
rng = random.Random(42)

EU_NAMES = {"Germany": "Deutschland", "Netherlands": "Nederland", "Spain": "España", "Switzerland": "Schweiz",
            "Belgium": "Belgique", "Austria": "Österreich", "Italy": "Italia", "Norway": "Norge",
            "Sweden": "Sverige", "Denmark": "Danmark", "Finland": "Suomi", "Poland": "Polska"}


def load_lines():
    con = duckdb.connect("dbt/warehouse.duckdb", read_only=True)
    return con.sql("""
        select invoice,
               row_number() over (partition by invoice order by source_sheet, source_row) as line,
               invoiced_at, stock_code, product_description, quantity, unit_price, customer_id, country, region
        from fct_retail_lines
        where month_start = date '2011-11-01'
        order by invoiced_at, invoice, line
    """).fetchall()


def money_text(v):
    return f"{v:.2f}" if v == round(v, 2) else f"{v:.3f}"


def uk_file(rows, preview):
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Sales")
    cells = []

    def put(values, issue=None):
        ws.append(values)
        cells.append({"values": ["" if v is None else str(v) for v in values], "issue": issue})

    put(["Sales export: UK team"], "Title rows above the real header")
    put(["Exported from the till system 02/12/2011"], "Title rows above the real header")
    put([])
    put(["Inv No", "Line", "Date", "Item code", "Item", "Qty", "Unit £", "Line total £", "Cust #", "Country"])
    by_day = {}
    for r in rows:
        by_day.setdefault(r[2].date(), []).append(r)
    dup_days = set(rng.sample(sorted(by_day), 3))
    for day, day_rows in by_day.items():
        block = []
        for r in day_rows:
            inv, line, ts, code, desc, qty, price, cust, country, _ = r
            issue = None
            date_val = ts
            if rng.random() < 0.03:
                date_val, issue = ts.strftime("%d/%m/%Y %H:%M"), "Date typed as text"
            price_val = float(price)
            if rng.random() < 0.05:
                price_val, issue = f"£{money_text(float(price))}", "Price typed as text with £"
            ctry = country
            roll = rng.random()
            if roll < 0.10:
                ctry, issue = "UK", "Country spelled differently"
            elif roll < 0.13:
                ctry, issue = "U.K.", "Country spelled differently"
            values = [inv, line, date_val, code, desc, qty, price_val, round(qty * float(price), 3),
                      int(cust) if cust else None, ctry]
            put(values, issue)
            block.append(values)
        if day in dup_days:  # someone pasted the last 40 rows of the day twice
            for values in block[-40:]:
                put(values, "Rows pasted twice")
        total = round(sum(qty * float(price) for _, _, _, _, _, qty, price, _, _, _ in day_rows), 2)
        put([f"Total {day:%d/%m/%Y}", None, None, None, None, None, None, total, None, None], "Subtotal row inside the data")
        put([], "Blank separator row")
    path = OUT / "UK_sales_Nov-2011.xlsx"
    wb.save(path)
    preview["files"].append({"name": path.name, "team": "UK team", "format": "Excel", "rows_in_file": len(cells),
                             "sample": pick_sample(cells)})


def europe_file(rows, preview):
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Verkauf")
    cells = []

    def put(values, issue=None):
        ws.append(values)
        cells.append({"values": ["" if v is None else str(v) for v in values], "issue": issue})

    put(["EUROPE TEAM — NOVEMBER 2011"], "Title rows above the real header")
    put(["Date", "Invoice", "Pos", "SKU", "Description", "Price", "Quantity", "Customer", "Country"])
    for r in rows:
        inv, line, ts, code, desc, qty, price, cust, country, _ = r
        issue = "Decimal commas and text dates"
        price_txt = money_text(float(price)).replace(".", ",")
        if float(price) >= 1000:
            whole, frac = price_txt.split(",")
            price_txt = f"{int(whole):,}".replace(",", ".") + "," + frac
            issue = "Thousands dot and decimal comma"
        qty_val = str(qty)
        if rng.random() < 0.10:
            qty_val, issue = f"{qty} pcs", "Quantity typed with units"
        sku = code
        if rng.random() < 0.10:
            sku, issue = code.lower() + " ", "Lower-case code with trailing space"
        ctry = EU_NAMES.get(country, country)
        if ctry != country and issue == "Decimal commas and text dates":
            issue = "Country in the local language"
        put([ts.strftime("%d.%m.%Y %H:%M"), inv, line, sku, desc, price_txt, qty_val,
             f"C-{cust}" if cust else "-", ctry], issue)
    path = OUT / "Europe_sales_Nov2011.xlsx"
    wb.save(path)
    preview["files"].append({"name": path.name, "team": "Europe team", "format": "Excel", "rows_in_file": len(cells),
                             "sample": pick_sample(cells)})


def row_file(rows, preview):
    path = OUT / "RestOfWorld_sales_Nov2011.csv"
    cells = []
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        header = ["invoice_no", "line", "invoice_datetime", "product_code", "product_name", "qty", "unit_price",
                  "customer_id", "country", "type"]
        w.writerow(header)
        cells.append({"values": header, "issue": None})
        for r in rows:
            inv, line, ts, code, desc, qty, price, cust, country, _ = r
            us_time = f"{ts.month}/{ts.day}/{ts.year} {(ts.hour % 12) or 12}:{ts.minute:02d} {'AM' if ts.hour < 12 else 'PM'}"
            kind = "Return" if qty < 0 else "Sale"
            alias = {"United States": "USA", "South Africa": "RSA", "South Korea": "Korea"}.get(country, country)
            values = [inv, line, us_time, f" {code}" if rng.random() < 0.05 else code, desc, abs(qty),
                      money_text(float(price)), cust or "N/A", alias, kind]
            w.writerow(values)
            issue = "Returns as positive qty + a Type column" if kind == "Return" else "US-style dates (month first, 12-hour)"
            cells.append({"values": [str(v) for v in values], "issue": issue})
    preview["files"].append({"name": path.name, "team": "Rest-of-world team", "format": "CSV", "rows_in_file": len(cells),
                             "sample": pick_sample(cells)})


def pick_sample(cells):
    """First rows of the file plus one example of every planted problem, for the 'before' preview."""
    sample = [dict(c, row=i + 1) for i, c in enumerate(cells[:10])]
    seen = {c["issue"] for c in sample}
    for i, c in enumerate(cells):
        if c["issue"] and c["issue"] not in seen:
            sample.append(dict(c, row=i + 1))
            seen.add(c["issue"])
    return sample


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    PREVIEW.parent.mkdir(parents=True, exist_ok=True)
    rows = load_lines()
    preview = {"source_lines": len(rows), "month": "November 2011", "files": []}
    uk_file([r for r in rows if r[9] == "United Kingdom"], preview)
    europe_file([r for r in rows if r[9] == "Europe"], preview)
    row_file([r for r in rows if r[9] == "Rest of world"], preview)
    for f in preview["files"]:
        f["bytes"] = (OUT / f["name"]).stat().st_size
    PREVIEW.write_text(json.dumps(preview, default=str))
    print(f"wrote {len(preview['files'])} messy files from {len(rows)} real lines")


if __name__ == "__main__":
    main()
